"""Regression pins for the corrected chlorine-capacity workbook."""

from __future__ import annotations

import re
import subprocess
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook
import pytest
import yaml

from simulator.battery.enums import Phase, PhaseRole, Quantity, ValueKind
from tests.battery.test_migrate import REPO_ROOT, _migrate_real_extract
from tests.helpers.corpus_paths import regolith_corpus_root


SOURCE_ID = "rusiecka-2026-chlorine-data"
EXTRACT_NAME = f"{SOURCE_ID}.yaml"
EXTRACT_PATH = REPO_ROOT / "data" / "literature" / "extracts" / EXTRACT_NAME
WORKBOOK_NAME = "rusiecka-2026-chlorine-data.xlsx"


def _workbook_or_skip() -> Path:
    workbook = regolith_corpus_root() / "raw" / SOURCE_ID / WORKBOOK_NAME
    if not workbook.is_file():
        pytest.skip(f"source workbook is absent: {workbook}")
    return workbook


def test_corrected_rows_migrate_and_pin_workbook_cells(tmp_path: Path) -> None:
    workbook = _workbook_or_skip()
    ws = load_workbook(workbook, data_only=True, read_only=False)["data"]
    doc = yaml.safe_load(EXTRACT_PATH.read_text(encoding="utf-8"))
    result = _migrate_real_extract(tmp_path / "migration", EXTRACT_NAME)
    assert result.validation is not None
    assert not result.validation.hard_issues, result.validation.hard_issues

    raw_rows = {
        observation["observation_id"]: (formula, observation)
        for formula, container in doc["species"].items()
        for observation in container.get("observations", [])
    }
    migrated = {
        row.observation_id.rsplit("::", 1)[-1]: row
        for row in result.observations.values()
    }
    capacities = [
        (formula, observation)
        for formula, observation in raw_rows.values()
        if observation["values"].get("quantity") == "chloride_capacity"
    ]
    compositions = [
        (formula, observation)
        for formula, observation in raw_rows.values()
        if observation["values"].get("quantity") == "equilibrium_phase_composition"
    ]
    expected_component_cells = sum(
        ws.cell(row, column).value is not None
        for row in range(81, 129)
        for column in range(2, 19)
    )
    assert len(capacities) == 126
    assert len(compositions) == expected_component_cells == 816
    assert len(raw_rows) == len(migrated) == 942

    for formula, observation in capacities:
        row_number = int(re.search(r"row(\d+)_chloride_capacity$", observation["observation_id"]).group(1))
        row = migrated[observation["observation_id"]]
        assert formula == row.identity.species.formula == "Cl"
        assert row.identity.quantity.value is Quantity.CHLORIDE_CAPACITY
        assert row.identity.phase_role.value is PhaseRole.MELT
        assert row.identity.species.phase.value is Phase.L
        assert row.value.kind is ValueKind.POINT
        assert not row.notices
        assert row.value.point == Decimal(f"{ws.cell(row_number, 28).value:.2f}")
        assert row.identity.temperature_K.is_value
        assert row.identity.temperature_K.value == Decimal(str(ws.cell(row_number, 20).value)) + Decimal("273.15")
        assert row.identity.total_pressure_Pa.is_unknown
        assert row.identity.fO2_Pa.is_unknown
        assert row.experiment_id.endswith(f"::experiment::worksheet-row-{row_number}")
        assert row.point_conditions["fCl2_bar"].state.is_value
        assert row.point_conditions["fCl2_bar"].locator is not None
        assert row.point_conditions["fO2_Pa"].state.is_value
        assert row.point_conditions["total_pressure_Pa"].state.is_value
        assert row.point_conditions["total_pressure_Pa"].locator is not None
        assert row.point_conditions["fO2_Pa"].locator is not None
        assert "fO2_log" not in row.point_conditions

    for formula, observation in compositions:
        match = re.search(r"cell ([A-Z]+)(\d+)$", observation["locator"]["record"])
        assert match is not None
        column_letters, row_text = match.groups()
        row_number = int(row_text)
        raw_id = observation["observation_id"]
        row = migrated[raw_id]
        source_cell = ws[f"{column_letters}{row_number}"]
        displayed_component = (
            Decimal(f"{source_cell.value:.2f}")
            if source_cell.number_format == "0.00"
            else Decimal(format(float(source_cell.value), ".15g"))
        )
        assert row.identity.species.formula == formula
        assert row.identity.quantity.value is Quantity.EQUILIBRIUM_PHASE_COMPOSITION
        assert row.identity.phase_role.value is PhaseRole.GLASS
        assert row.identity.species.phase.value is Phase.GLASS
        assert row.identity.subtype.value == "wt_percent"
        assert row.identity.temperature_K.value == Decimal(str(ws.cell(row_number, 20).value)) + Decimal("273.15")
        assert row.value.kind is ValueKind.POINT
        assert row.value.point == displayed_component
        assert row.identity.total_pressure_Pa.is_unknown
        assert row.identity.fO2_Pa.is_unknown
        assert row.experiment_id.endswith(f"::experiment::worksheet-row-{row_number}")
        assert row.point_conditions["total_pressure_Pa"].state.is_value
        assert row.point_conditions["fO2_Pa"].state.is_value
        assert row.point_conditions["fCl2_bar"].state.is_value
        assert not row.notices

    # Independent cell pins: Lunar-1 is a named legacy composition row, while AND-1 is this study.
    assert ws["AC22"].value == "Lunar-1"
    assert ws["B22"].value == 45.77
    assert ws["K22"].value == 2.03
    assert ws["T22"].value == 1400
    assert ws["U22"].value == 1.63
    assert round(ws["AB22"].value, 2) == 0.20
    assert "rusiecka_2026_data_row22_sio2" not in migrated
    assert "rusiecka_2026_data_row22_cl" not in migrated
    assert migrated["rusiecka_2026_data_row22_chloride_capacity"].value.point == Decimal("0.20")

    assert ws["AC81"].value == "AND-1"
    assert ws["T81"].value == 1200
    assert ws["U81"].value == 1.05
    assert f"{ws['X81'].value:.2E}" == "2.80E-05"
    assert f"{ws['Y81'].value:.2E}" == "5.54E-10"
    assert round(ws["AA81"].value, 2) == 1.68
    assert round(ws["AB81"].value, 2) == 0.23
    and1 = migrated["rusiecka_2026_data_row81_chloride_capacity"]
    assert and1.value.point == Decimal("0.23")
    assert and1.point_conditions["total_pressure_Pa"].state.value == Decimal("1050000000.00")
    assert and1.point_conditions["fO2_Pa"].state.value == Decimal("0.000055400000000")


def test_companion_table2_locator_and_sample_list_disagreement() -> None:
    workbook = _workbook_or_skip()
    corpus = regolith_corpus_root()
    pdf = corpus / "raw" / "rusiecka-wood-2026-chlorine-hcl" / "rusiecka-wood-2026-chlorine-hcl.pdf"
    if not pdf.is_file():
        pytest.skip(f"companion paper PDF is absent: {pdf}")
    ws = load_workbook(workbook, data_only=True, read_only=False)["data"]
    result = subprocess.run(
        ["pdftotext", "-f", "5", "-l", "5", "-layout", str(pdf), "-"],
        check=True,
        capture_output=True,
        text=True,
    )
    page = re.sub(r"\s+", " ", result.stdout)
    worksheet_ids = {ws.cell(row, 29).value for row in range(81, 97)}
    assert "AND-7" not in worksheet_ids
    assert "DCB-6" in worksheet_ids
    assert "AND-7" in page
    assert "DCB-6" not in page
    assert "6.08E-06" in page
    assert "4.44E-11" in page
