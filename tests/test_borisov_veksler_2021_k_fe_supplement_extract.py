"""Migration pins for Borisov and Veksler (2021), supplementary workbook."""

from collections import Counter
from decimal import Decimal
from pathlib import Path
import re

import pytest
import yaml
from openpyxl import load_workbook

from tests.battery.factories import migrate_one_extract
from tests.helpers.corpus_paths import regolith_corpus_root


ROOT = Path(__file__).resolve().parents[1]
SOURCE_ID = "borisov-veksler-2021-k-fe-supplement"
EXTRACT = ROOT / f"data/literature/extracts/{SOURCE_ID}.yaml"
SOURCE_XLSX = (
    regolith_corpus_root()
    / "raw"
    / SOURCE_ID
    / f"{SOURCE_ID}.xlsx"
)
def _displayed_cell(cell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        match = re.fullmatch(r"0(?:\.(0+))?", cell.number_format)
        if match and match.group(1) is not None:
            return format(value, f".{len(match.group(1))}f")
        return format(value, ".15g")
    return str(value).strip()


@pytest.fixture(scope="module")
def migrated_extract():
    document = yaml.safe_load(EXTRACT.read_text(encoding="utf-8"))
    migrator = migrate_one_extract(ROOT, EXTRACT, index={}, aliases={})
    migrator.finalize()
    return document, migrator.result


@pytest.mark.skipif(not SOURCE_XLSX.is_file(), reason="regolith-corpus workbook is absent")
def test_every_workbook_component_migrates_as_its_printed_point_or_refusal(migrated_extract) -> None:
    document, result = migrated_extract
    assert not result.validation.hard_issues
    observations = document["species"]["SiO2"]["observations"]
    assert len(observations) == 10_092

    workbook = load_workbook(SOURCE_XLSX, data_only=True, read_only=False)
    row_counts = Counter()
    numeric = refusals = 0
    for raw in observations:
        values = raw["values"]
        locator = raw["locator"]
        sheet = locator["table"]
        source_cell = workbook[sheet][locator["record"]]
        row = values["source_row_as_published"]
        formula_as_printed = values["formula_as_published"]
        headers = {
            str(workbook[sheet].cell(3, col).value).strip(): col
            for col in range(1, workbook[sheet].max_column + 1)
            if workbook[sheet].cell(3, col).value is not None
        }
        assert source_cell.row == row
        assert source_cell.value == workbook[sheet].cell(row, headers[formula_as_printed]).value
        assert values["formula"] == ("FeO" if formula_as_printed == "FeOt" else formula_as_printed)
        row_counts[sheet] += 1

        observation = result.observations[f"{SOURCE_ID}::{raw['observation_id']}"]
        assert observation.identity.quantity.value.value == "equilibrium_phase_composition"
        assert observation.identity.species.formula == values["formula"]
        assert observation.identity.species.phase.value.value == "l"
        assert observation.identity.phase_role.value.value == "melt"
        assert observation.identity.subtype.value == "wt_percent"
        assert observation.point_conditions is not None

        printed = _displayed_cell(source_cell)
        if isinstance(source_cell.value, (int, float)) and not isinstance(source_cell.value, bool):
            expected = Decimal(printed)
            assert values["admission_status"] == "admitted"
            assert observation.value.kind.value == "point"
            assert observation.value.point == expected
            assert observation.notices == ()
            numeric += 1
        else:
            assert values["admission_status"] == "typed_refusal"
            assert values["printed_cell_as_published"] == printed
            assert "amount is unknown" in values["reason"] or "not a numeric point" in values["reason"]
            assert "component_amount" not in values
            assert observation.value.kind.value == "unavailable"
            refusals += 1

        total_col = headers["Total"]
        total_cell = workbook[sheet].cell(row, total_col)
        total_condition = observation.point_conditions["printed_sum"]
        assert total_condition.state.value == Decimal(_displayed_cell(total_cell))
        assert total_condition.locator.table == sheet
        assert total_condition.locator.record == total_cell.coordinate

        if sheet == "natural immiscible liquids":
            assert observation.identity.temperature_K.is_unknown
            assert observation.identity.total_pressure_Pa.is_unknown
            assert observation.point_conditions["temperature_K"].state.is_unknown
            assert observation.point_conditions["total_pressure_Pa"].state.is_unknown
            continue

        temperature_col = headers["T, C"]
        temperature_cell = workbook[sheet].cell(row, temperature_col)
        printed_temperature = Decimal(_displayed_cell(temperature_cell))
        assert observation.identity.temperature_K.value == printed_temperature + Decimal("273.15")
        t_condition = observation.point_conditions["temperature_C_as_published"]
        assert t_condition.state.value == printed_temperature
        assert t_condition.locator.record == temperature_cell.coordinate

        pressure_col = headers["P, kbar"]
        pressure_cell = workbook[sheet].cell(row, pressure_col)
        printed_pressure = Decimal(_displayed_cell(pressure_cell))
        assert observation.identity.total_pressure_Pa.value == printed_pressure * Decimal("100000000")
        p_condition = observation.point_conditions["pressure_as_published"]
        assert p_condition.state.value == printed_pressure
        assert p_condition.locator.record == pressure_cell.coordinate
        assert values["total_pressure_Pa"] == {
            "value": _displayed_cell(pressure_cell),
            "units": "kbar",
        }

        oxygen_col = headers["logfO2"]
        oxygen_cell = workbook[sheet].cell(row, oxygen_col)
        oxygen_text = _displayed_cell(oxygen_cell)
        if oxygen_text:
            try:
                printed_log = Decimal(oxygen_text)
            except (ArithmeticError, ValueError):
                if oxygen_text.strip().casefold() == "in ar":
                    condition = observation.point_conditions["atmosphere"]
                else:
                    condition = observation.point_conditions["fO2_control"]
                assert condition.state.value == oxygen_text
                assert condition.locator.record == oxygen_cell.coordinate
            else:
                log_condition = observation.point_conditions["fO2_log"]
                assert log_condition.state.value == printed_log
                assert log_condition.locator.record == oxygen_cell.coordinate
                assert observation.identity.fO2_Pa.is_unknown
        else:
            assert not any(key in observation.point_conditions for key in ("fO2_log", "fO2_control", "atmosphere"))

    assert row_counts == Counter({
        "included experimental data": 4_656,
        "excluded experimental data": 4_316,
        "natural immiscible liquids": 1_120,
    })
    assert numeric == 8_737
    assert refusals == 1_355
    assert numeric + refusals == len(observations)
    assert result.observations and len(result.observations) == len(observations)


@pytest.mark.skipif(not SOURCE_XLSX.is_file(), reason="regolith-corpus workbook is absent")
def test_independent_cell_pins_fail_when_source_or_extract_value_is_mutated(migrated_extract) -> None:
    document, _result = migrated_extract
    workbook = load_workbook(SOURCE_XLSX, data_only=True, read_only=False)
    pins = (
        ("included experimental data", "G4", "SiO2", "55.8"),
        ("excluded experimental data", "F4", "SiO2", "49.20"),
        ("natural immiscible liquids", "C4", "SiO2", "40.8"),
        ("included experimental data", "J4", "FeOt", "40.1"),
    )
    raw_by_cell = {
        (record["locator"]["table"], record["locator"]["record"]): record
        for record in document["species"]["SiO2"]["observations"]
    }

    def assert_pin(source_value: str, extract_value: str, expected: str) -> None:
        assert source_value == expected
        assert extract_value == expected

    for sheet, cell_ref, _component, expected in pins:
        source_value = _displayed_cell(workbook[sheet][cell_ref])
        record = raw_by_cell[(sheet, cell_ref)]
        extract_value = record["values"]["component_amount"]["value"]
        assert_pin(source_value, extract_value, expected)
        with pytest.raises(AssertionError):
            assert_pin(str(Decimal(source_value) + Decimal("1")), extract_value, expected)
        with pytest.raises(AssertionError):
            assert_pin(source_value, str(Decimal(extract_value) + Decimal("1")), expected)


@pytest.mark.skipif(not SOURCE_XLSX.is_file(), reason="regolith-corpus workbook is absent")
def test_context_preserves_unlabelled_and_formula_cells(migrated_extract) -> None:
    document, _result = migrated_extract
    context = document["species"]["SiO2"]["context"]
    by_id = {item["observation_id"]: item for item in context}
    unlabelled = by_id[f"{SOURCE_ID.replace('-', '_')}_sheet1_unlabelled_context"]["values"]
    assert unlabelled["cell_count"] == 350
    assert len(unlabelled["cells_as_published"]) == 350
    formula = by_id[f"{SOURCE_ID.replace('-', '_')}_excluded_sheet_formula_context"]["values"]
    assert len(formula["formula_rows_as_published"]) == 22
    assert all(
        cell["formula_as_stored"].startswith("=")
        for row in formula["formula_rows_as_published"]
        for cell in row["cells"]
    )
